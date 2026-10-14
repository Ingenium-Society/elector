"""Export the final election results as an official, printable report.

Run this ONCE, from the offline/ folder, after the admin has pressed
"Reveal results" (state RESULTS):

    python scripts/export_results.py
    python scripts/export_results.py --roll-ticks 118     # ticks counted on the paper voter roll

It creates a folder (default: data/exports/) with:

    results_report.docx      the official report: totals, reconciliation, SHA-256, signature lines
    results.xlsx             candidate name and total votes only
    results.csv              the same, as plain text
    election.sqlite          a consistent copy of the final database
    election.sqlite.sha256   SHA-256 hash of that copy, so anyone can verify it later

ANONYMITY (AGENTS.md rule 8, SPEC A10): this script only ever asks the database for
COUNTS (votes per candidate, tokens used / unused). It never joins the tokens table
with the votes table, and it never writes a token or a token hash into a report,
a spreadsheet or a CSV. The exports folder is inside data/, which git ignores.
"""
import argparse
import csv
import hashlib
import sqlite3
import sys
from datetime import datetime
from pathlib import Path

# Let Python find the "app" package, which lives one folder above scripts/.
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from app import config, db  # noqa: E402

LOGO = Path(__file__).resolve().parent.parent / "app" / "static" / "logo.png"
SCHOOL = "Global Darussalam Academy"

REPORT_NAME = "results_report.docx"
XLSX_NAME = "results.xlsx"
CSV_NAME = "results.csv"
DB_COPY_NAME = "election.sqlite"
HASH_NAME = DB_COPY_NAME + ".sha256"
ALL_OUTPUTS = (REPORT_NAME, XLSX_NAME, CSV_NAME, DB_COPY_NAME, HASH_NAME)

NAVY = "0F2268"
GOLD = "F6D463"
GOLD_SOFT = "FFF6D6"
LIGHT = "EEF3FB"

PASS, FAIL, TODO = "PASS", "FAIL", "TO COMPLETE"


# ---- 1. Read the numbers (counts only) --------------------------------------

def collect():
    """Return the aggregate numbers. Nothing here links a token to a vote."""
    candidates = db.result_totals()  # [{"id", "name", "votes"}, ...] in id order
    with db.connect() as conn:
        votes_rows = conn.execute("SELECT COUNT(*) FROM votes").fetchone()[0]
        issued = conn.execute("SELECT COUNT(*) FROM tokens").fetchone()[0]
        used = conn.execute("SELECT COUNT(*) FROM tokens WHERE used = 1").fetchone()[0]
        unused = conn.execute("SELECT COUNT(*) FROM tokens WHERE used = 0").fetchone()[0]
    return {
        "candidates": candidates,
        "votes_rows": votes_rows,
        "issued": issued,
        "used": used,
        "unused": unused,
    }


def percent(part, whole):
    return "0.0%" if whole == 0 else f"{100 * part / whole:.1f}%"


def leaders(candidates):
    """Candidates with the highest number of votes (more than one means a tie)."""
    top = max((c["votes"] for c in candidates), default=0)
    if top == 0:
        return []
    return [c for c in candidates if c["votes"] == top]


# ---- 2. Reconciliation (SPEC section 11) ------------------------------------

def reconcile(data, roll_ticks=None):
    """Return a list of (check, detail, status). Status is PASS, FAIL or TO COMPLETE."""
    issued, used, unused = data["issued"], data["used"], data["unused"]
    votes = data["votes_rows"]
    counted = sum(c["votes"] for c in data["candidates"])

    checks = [
        (
            "Issued = used + unused tokens",
            f"{issued} = {used} + {unused}",
            PASS if issued == used + unused else FAIL,
        ),
        (
            "Tokens used = ballots stored",
            f"{used} tokens used, {votes} ballots stored",
            PASS if used == votes else FAIL,
        ),
        (
            "Candidate totals = ballots stored",
            f"{counted} counted for candidates, {votes} ballots stored",
            PASS if counted == votes else FAIL,
        ),
    ]
    if roll_ticks is None:
        checks.append((
            "Voter roll ticks <= tokens issued",
            f"Committee to enter the ticks counted on the paper roll (at most {issued})",
            TODO,
        ))
    else:
        checks.append((
            "Voter roll ticks <= tokens issued",
            f"{roll_ticks} ticks on the paper roll, {issued} tokens issued",
            PASS if roll_ticks <= issued else FAIL,
        ))
    return checks


# ---- 3. Database copy and its SHA-256 ---------------------------------------

def snapshot_database(destination):
    """Copy the live database with SQLite's own backup, which is safe while the server runs."""
    source = sqlite3.connect(config.DB_PATH)
    target = sqlite3.connect(destination)
    try:
        source.backup(target)
    finally:
        target.close()
        source.close()


def sha256_of(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


# ---- 4. Spreadsheet and CSV (candidate name and votes only) ------------------

def write_csv(path, candidates):
    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["Candidate", "Votes"])
        for c in candidates:
            writer.writerow([c["name"], c["votes"]])


def write_xlsx(path, candidates):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "Results"
    sheet.append(["Candidate", "Votes"])
    for c in candidates:
        sheet.append([c["name"], c["votes"]])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
    for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].alignment = Alignment(horizontal="right")
    sheet.column_dimensions["A"].width = 36
    sheet.column_dimensions["B"].width = 10
    book.save(path)


# ---- 5. The official report (.docx) -----------------------------------------

def write_report(path, data, checks, digest, roll_ticks):
    """Build the printable report. The helpers below only style things (colours, borders)."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    def rgb(hex_color):
        return RGBColor.from_string(hex_color)

    def shade(cell, fill):
        props = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        props.append(shd)

    def line_under(paragraph, color, size=8):
        """Draw a line under a paragraph (used for the gold rule and signature lines)."""
        p_pr = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), str(size))
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), color)
        borders.append(bottom)
        p_pr.append(borders)

    def light_borders(table, color="C9D2E3"):
        props = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            element = OxmlElement(f"w:{edge}")
            element.set(qn("w:val"), "single")
            element.set(qn("w:sz"), "4")
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), color)
            borders.append(element)
        props.append(borders)

    def text_in(cell, text, bold=False, color=None, size=10.5, align=None, font=None):
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(3)
        paragraph.paragraph_format.space_after = Pt(3)
        if align:
            paragraph.alignment = align
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = rgb(color)
        if font:
            run.font.name = font
        return paragraph

    def make_table(widths, borders=True):
        table = doc.add_table(rows=0, cols=len(widths))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        # Word and LibreOffice both need the column widths on the table itself,
        # not only on each cell, otherwise every column gets the same width.
        for i, width in enumerate(widths):
            table.columns[i].width = Cm(width)
        if borders:
            light_borders(table)
        return table

    def add_row(table, widths, values, bold=False, fill=None, color=None, aligns=None):
        row = table.add_row()
        for i, cell in enumerate(row.cells):
            cell.width = Cm(widths[i])
            if fill:
                shade(cell, fill)
            align = aligns[i] if aligns else None
            text_in(cell, str(values[i]), bold=bold, color=color, align=align)
        return row

    def heading(text):
        paragraph = doc.add_heading(text, level=2)
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(4)
        return paragraph

    doc = Document()
    doc.core_properties.title = "Official Election Results"
    doc.core_properties.author = "School Election System"

    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)  # A4
    section.left_margin = section.right_margin = Cm(2)
    section.top_margin, section.bottom_margin = Cm(1.3), Cm(1.3)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    h2 = doc.styles["Heading 2"]
    h2.font.name = "Cambria"
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = rgb(NAVY)
    h2.element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
    h2.element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")

    # -- Title block with the school logo ------------------------------------
    top = doc.add_paragraph()
    top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    top.paragraph_format.space_after = Pt(2)
    if LOGO.exists():
        top.add_run().add_picture(str(LOGO), width=Cm(2.3))

    school = doc.add_paragraph()
    school.alignment = WD_ALIGN_PARAGRAPH.CENTER
    school.paragraph_format.space_after = Pt(0)
    run = school.add_run(SCHOOL.upper())
    run.bold = True
    run.font.size = Pt(10)
    run.font.name = "Cambria"
    run.font.color.rgb = rgb(NAVY)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(0)
    run = title.add_run("Official Election Results")
    run.bold = True
    run.font.size = Pt(24)
    run.font.name = "Cambria"
    run.font.color.rgb = rgb(NAVY)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(6)
    stamp = datetime.now().astimezone().strftime("%d %B %Y, %H:%M")
    run = sub.add_run(f"Student council election  |  Report generated {stamp}")
    run.font.size = Pt(9.5)
    run.font.color.rgb = rgb("526173")
    line_under(sub, GOLD, size=18)

    # -- Result banner --------------------------------------------------------
    candidates = data["candidates"]
    votes_total = data["votes_rows"]
    top_candidates = leaders(candidates)
    if not top_candidates:
        banner = "No votes were cast."
    elif len(top_candidates) == 1:
        c = top_candidates[0]
        banner = f"{c['name']} received the most votes: {c['votes']} of {votes_total} ({percent(c['votes'], votes_total)})."
    else:
        names = " and ".join(c["name"] for c in top_candidates)
        banner = (f"TIE: {names} each received {top_candidates[0]['votes']} votes. "
                  "The election committee must apply its agreed tie-break rule.")

    box = make_table([17])
    row = box.add_row()
    row.cells[0].width = Cm(17)
    shade(row.cells[0], GOLD_SOFT)
    text_in(row.cells[0], banner, bold=True, color=NAVY, size=12.5, align=WD_ALIGN_PARAGRAPH.CENTER)

    # -- Vote totals ----------------------------------------------------------
    heading("Vote totals")
    widths = [1.6, 9.4, 3.0, 3.0]
    right, center = WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER
    table = make_table(widths)
    add_row(table, widths, ["Rank", "Candidate", "Votes", "Share"], bold=True, fill=NAVY,
            color="FFFFFF", aligns=[center, None, right, right])
    ranked = sorted(candidates, key=lambda c: (-c["votes"], c["id"]))
    previous_votes, rank = None, 0
    for position, c in enumerate(ranked, start=1):
        if c["votes"] != previous_votes:
            rank = position  # equal votes share the same rank
        previous_votes = c["votes"]
        is_leader = c in top_candidates
        add_row(table, widths, [rank, c["name"], c["votes"], percent(c["votes"], votes_total)],
                bold=is_leader, fill=GOLD_SOFT if is_leader else None, aligns=[center, None, right, right])
    add_row(table, widths, ["", "Total ballots", votes_total, "100.0%" if votes_total else "0.0%"],
            bold=True, fill=LIGHT, aligns=[center, None, right, right])

    # -- Turnout --------------------------------------------------------------
    heading("Turnout")
    widths = [4.25, 4.25, 4.25, 4.25]
    center = WD_ALIGN_PARAGRAPH.CENTER
    table = make_table(widths)
    add_row(table, widths, ["Tokens issued", "Ballots cast", "Tokens not used", "Turnout"],
            bold=True, fill=LIGHT, color=NAVY, aligns=[center] * 4)
    add_row(table, widths, [data["issued"], data["used"], data["unused"], percent(data["used"], data["issued"])],
            aligns=[center] * 4)

    # -- Reconciliation -------------------------------------------------------
    heading("Reconciliation checks")
    widths = [7.6, 7.2, 2.2]
    table = make_table(widths)
    add_row(table, widths, ["Check", "Numbers", "Result"], bold=True, fill=NAVY, color="FFFFFF",
            aligns=[None, None, center])
    colors = {PASS: "11643A", FAIL: "B42318", TODO: "8A6D00"}
    for label, detail, status in checks:
        row = add_row(table, widths, [label, detail, status], aligns=[None, None, center])
        for run in row.cells[2].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = rgb(colors[status])
    if any(status == FAIL for _, _, status in checks):
        warning = doc.add_paragraph()
        warning.paragraph_format.space_before = Pt(6)
        run = warning.add_run("RECONCILIATION FAILED. The observers must be told before these results are accepted.")
        run.bold = True
        run.font.color.rgb = rgb("B42318")

    # -- Integrity ------------------------------------------------------------
    heading("Data integrity")
    para = doc.add_paragraph(f"SHA-256 of the final database file ({DB_COPY_NAME}):")
    para.paragraph_format.space_after = Pt(2)
    code = doc.add_paragraph()
    code.paragraph_format.space_after = Pt(4)
    run = code.add_run(digest)
    run.font.name = "Courier New"
    run.font.size = Pt(9.5)
    run.bold = True
    note = doc.add_paragraph()
    note.paragraph_format.space_after = Pt(0)
    run = note.add_run(
        "To check later that the file was not changed: on Windows run  certutil -hashfile election.sqlite SHA256 ; "
        "on macOS or Linux run  shasum -a 256 election.sqlite . The result must match the value above."
    )
    run.font.size = Pt(9)
    run.font.color.rgb = rgb("526173")

    # -- Signatures -----------------------------------------------------------
    heading("Confirmation")
    intro = doc.add_paragraph("We confirm that the totals above were read from the election system in our presence.")
    intro.paragraph_format.space_after = Pt(2)
    roles = ["Election committee chair", "Observer (Class Rep. Council)", "Teacher / IT reviewer"]
    widths = [5.66, 5.66, 5.66]
    sign = make_table(widths, borders=False)
    role_row = sign.add_row()
    for i, role in enumerate(roles):
        role_row.cells[i].width = Cm(widths[i])
        text_in(role_row.cells[i], role, bold=True, color=NAVY, size=10)
    # One table row per line, so the three signature columns always line up and each
    # line is drawn (neighbouring paragraphs with the same border would merge into one).
    for label, space in (("Signature", 24), ("Name", 10), ("Date", 10)):
        line_row = sign.add_row()
        for i in range(3):
            cell = line_row.cells[i]
            cell.width = Cm(widths[i])
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(space)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.right_indent = Cm(0.5)
            run = p.add_run(label)
            run.font.size = Pt(8.5)
            run.font.color.rgb = rgb("526173")
            line_under(p, "7A869A", size=6)

    # -- Footer ---------------------------------------------------------------
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("This report contains vote totals only. It holds no token, voter or individual ballot information.")
    run.font.size = Pt(8)
    run.font.color.rgb = rgb("526173")

    doc.save(path)


# ---- 6. Put it all together --------------------------------------------------

def main(argv):
    parser = argparse.ArgumentParser(description="Export the final election results.")
    parser.add_argument("--out", type=Path, default=config.DB_PATH.parent / "exports",
                        help="folder for the exported files (default: data/exports)")
    parser.add_argument("--roll-ticks", type=int, default=None,
                        help="number of ticks counted on the paper voter roll")
    args = parser.parse_args(argv)

    if not config.DB_PATH.exists():
        sys.exit(f"Database not found: {config.DB_PATH}")
    state = db.get_state()
    if state != "RESULTS":
        # Totals must not leave the system before the admin has revealed them.
        sys.exit(f"The election state is {state}. Export is only allowed after 'Reveal results' (RESULTS).")

    taken = [name for name in ALL_OUTPUTS if (args.out / name).exists()]
    if taken:
        sys.exit(f"{args.out} already contains {', '.join(taken)}. Nothing was changed. "
                 "Move or delete that folder first if you really want a new export.")
    args.out.mkdir(parents=True, exist_ok=True)

    data = collect()
    checks = reconcile(data, args.roll_ticks)

    snapshot_database(args.out / DB_COPY_NAME)
    digest = sha256_of(args.out / DB_COPY_NAME)
    (args.out / HASH_NAME).write_text(f"{digest}  {DB_COPY_NAME}\n", encoding="utf-8")

    write_csv(args.out / CSV_NAME, data["candidates"])
    write_xlsx(args.out / XLSX_NAME, data["candidates"])
    write_report(args.out / REPORT_NAME, data, checks, digest, args.roll_ticks)

    print(f"Exported to {args.out}:")
    for name in ALL_OUTPUTS:
        print(f"  {name}")
    print(f"SHA-256 of {DB_COPY_NAME}: {digest}")
    for label, detail, status in checks:
        print(f"  [{status}] {label}: {detail}")
    if any(status == FAIL for _, _, status in checks):
        print("RECONCILIATION FAILED: tell the observers before accepting the results.")
        return 2
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
